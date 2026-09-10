"""Document text extraction for SIFguard.

Supported: PDF (native text via PyMuPDF, OCR fallback for scanned pages),
DOCX, TXT, and images (JPG/JPEG/PNG/WEBP via EasyOCR).

Notes:
  * The EasyOCR reader is created lazily on first use — importing this
    module must stay cheap so `/health` works even before OCR weights load.
  * Legacy `.doc` (OLE) files are NOT decodable without antiword/LibreOffice;
    they raise a clear error instead of returning garbage.
  * Corrupt/empty files raise ValueError with a human-readable message so
    the API can return 422 instead of crashing.
"""

from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
}

_reader = None


def _get_reader():
    global _reader
    if _reader is None:
        import easyocr

        _reader = easyocr.Reader(["en"])
    return _reader


def _read_upload_bytes(file_path):
    path = Path(file_path)
    if not path.is_file():
        raise ValueError(f"Uploaded file not found: {path.name}")
    if path.stat().st_size == 0:
        raise ValueError(f"Uploaded file is empty: {path.name}")
    return path


def extract_text_from_pdf(file_path):
    import fitz

    _read_upload_bytes(file_path)
    try:
        doc = fitz.open(file_path)
    except Exception as error:
        raise ValueError(f"Corrupt or unreadable PDF: {error}")

    try:
        text = ""
        for page in doc:
            text += page.get_text() + "\n"
    finally:
        doc.close()

    if len(text.strip()) >= 50:
        return text.strip()

    return extract_pdf_with_ocr(file_path)


def extract_pdf_with_ocr(file_path):
    import fitz

    reader = _get_reader()
    try:
        doc = fitz.open(file_path)
    except Exception as error:
        raise ValueError(f"Corrupt or unreadable PDF: {error}")

    text = ""
    try:
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            image_path = f"{file_path}_page_{page.number}.png"
            pix.save(image_path)
            try:
                result = reader.readtext(image_path)
                for detection in result:
                    text += detection[1] + " "
            finally:
                Path(image_path).unlink(missing_ok=True)
            text += "\n"
    finally:
        doc.close()

    return text.strip()


def extract_text_from_image(file_path):
    _read_upload_bytes(file_path)
    reader = _get_reader()
    try:
        result = reader.readtext(str(file_path))
    except Exception as error:
        raise ValueError(f"Unreadable image file: {error}")

    return "\n".join(
        detection[1] for detection in result if len(detection) > 1
    ).strip()


def extract_text_from_docx(file_path):
    _read_upload_bytes(file_path)
    try:
        import docx
    except ImportError as error:
        raise ValueError(f"DOCX support is not installed: {error}")
    try:
        document = docx.Document(str(file_path))
    except Exception as error:
        raise ValueError(f"Corrupt or unreadable DOCX: {error}")

    paragraphs = [para.text for para in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)
    return "\n".join(paragraphs).strip()


def extract_text_from_txt(file_path):
    _read_upload_bytes(file_path)
    with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
        return handle.read().strip()


def extract_text(file_path, filename):
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    if extension in [".jpg", ".jpeg", ".png", ".webp"]:
        return extract_text_from_image(file_path)

    if extension == ".txt":
        return extract_text_from_txt(file_path)

    if extension == ".doc":
        raise ValueError(
            "Legacy .doc files are not supported by the NLP service. "
            "Please re-save the report as DOCX, PDF or TXT and retry."
        )

    raise ValueError(f"Unsupported file type: {extension or '(none)'}")
