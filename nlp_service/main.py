import io
import logging
import os
import re
from contextlib import asynccontextmanager
from typing import List, Optional

import torch
import easyocr

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from pydantic import BaseModel, Field

from PIL import Image
from pypdf import PdfReader
from docx import Document as DocxDocument

from transformers import (
    AutoTokenizer,
    AutoModelForTokenClassification,
)

import risk_engine


# ============================================================
# CONFIGURATION
# ============================================================

MODEL_PATH = "./models/sif-distilbert-ner"

# DistilBERT maximum token length
MAX_LENGTH = 512

# Number of overlapping tokens between chunks
CHUNK_STRIDE = 64

# Maximum uploaded file size
MAX_FILE_SIZE = 10 * 1024 * 1024

# Supported files
SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
}


# ============================================================
# LOGGING
# ============================================================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
)

logger = logging.getLogger("SIF")


# ============================================================
# GLOBAL MODELS
# ============================================================

ocr_reader = None
tokenizer = None
ner_model = None


# ============================================================
# REQUEST MODELS
# ============================================================

class RiskAssessmentRequest(BaseModel):

    hazards: List[str] = Field(
        default_factory=list
    )

    activities: List[str] = Field(
        default_factory=list
    )

    locations: List[str] = Field(
        default_factory=list
    )

    barrier_failures: List[str] = Field(
        default_factory=list
    )


# ============================================================
# APPLICATION LIFESPAN
# ============================================================

@asynccontextmanager
async def lifespan(app: FastAPI):

    global ocr_reader
    global tokenizer
    global ner_model

    logger.info("========================================")
    logger.info("Starting SIF AI Processing Service")
    logger.info("========================================")

    try:

        # ----------------------------------------------------
        # Load EasyOCR
        # ----------------------------------------------------

        logger.info("Loading EasyOCR...")

        ocr_reader = easyocr.Reader(
            ["en"],
            gpu=False
        )

        logger.info(
            "EasyOCR loaded successfully."
        )


        # ----------------------------------------------------
        # Check fine-tuned model
        # ----------------------------------------------------

        if not os.path.exists(MODEL_PATH):

            raise FileNotFoundError(
                f"Fine-tuned DistilBERT model not found at: "
                f"{MODEL_PATH}"
            )


        # ----------------------------------------------------
        # Load tokenizer
        # ----------------------------------------------------

        logger.info(
            "Loading DistilBERT tokenizer..."
        )

        tokenizer = AutoTokenizer.from_pretrained(
            MODEL_PATH
        )


        # ----------------------------------------------------
        # Load fine-tuned NER model
        # ----------------------------------------------------

        logger.info(
            "Loading fine-tuned DistilBERT NER model..."
        )

        ner_model = (
            AutoModelForTokenClassification
            .from_pretrained(MODEL_PATH)
        )

        ner_model.eval()


        # ----------------------------------------------------
        # Show labels
        # ----------------------------------------------------

        logger.info(
            "Loaded entity labels: %s",
            ner_model.config.id2label
        )

        logger.info(
            "SIF AI Service Ready"
        )


    except Exception as error:

        logger.exception(
            "Failed to load AI models: %s",
            error
        )

        raise


    yield


    logger.info(
        "Shutting down SIF AI service..."
    )


# ============================================================
# FASTAPI APPLICATION
# ============================================================

app = FastAPI(

    title="SIF Safety Intelligence Engine",

    description=(
        "Safety Intelligence Framework AI service "
        "for document parsing, OCR, NER, "
        "SIF precursor severity classification "
        "and risk assessment."
    ),

    version="1.0.0",

    lifespan=lifespan,
)


# ============================================================
# ROOT ENDPOINT
# ============================================================

@app.get("/")
async def root():

    return {

        "service": (
            "SIF Safety Intelligence Engine"
        ),

        "status": "running",

        "model": (
            "Fine-tuned DistilBERT NER"
        ),

        "severity_classifier": (
            "SIF Precursor Severity 1-5"
        ),

        "version": "1.0.0",

    }


# ============================================================
# HEALTH ENDPOINT
# ============================================================

@app.get("/health")
async def health_check():

    models_loaded = (

        ocr_reader is not None

        and tokenizer is not None

        and ner_model is not None

    )


    return {

        "status": (
            "healthy"
            if models_loaded
            else "starting"
        ),

        "ocr_loaded": (
            ocr_reader is not None
        ),

        "distilbert_loaded": (
            tokenizer is not None
            and ner_model is not None
        ),

        "severity_classifier": (
            "rule-based 1-5"
        ),

    }


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_report_text(
    text: str
) -> str:

    if not text:

        return ""


    # Normalize line endings

    text = (
        str(text)
        .replace("\r\n", "\n")
        .replace("\r", "\n")
    )


    # Remove control characters

    text = re.sub(
        r"[\x00-\x08\x0B\x0C\x0E-\x1F]",
        "",
        text
    )


    # Normalize spaces

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )


    # Remove excessive new lines

    text = re.sub(
        r"\n+",
        "\n",
        text
    )


    # Strip every line

    lines = [

        line.strip()

        for line in text.split("\n")

        if line.strip()

    ]


    text = "\n".join(
        lines
    )


    # Final whitespace cleanup

    text = re.sub(
        r"\s+",
        " ",
        text
    )


    return text.strip()


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

async def extract_text_from_pdf(
    document: UploadFile
) -> str:

    try:

        file_bytes = await document.read()


        if not file_bytes:

            raise ValueError(
                "PDF file is empty."
            )


        reader = PdfReader(
            io.BytesIO(file_bytes)
        )


        pages = []


        for page in reader.pages:

            page_text = page.extract_text()


            if page_text:

                pages.append(
                    page_text
                )


        return "\n".join(
            pages
        ).strip()


    except Exception as error:

        logger.exception(
            "PDF extraction failed: %s",
            error
        )


        raise HTTPException(

            status_code=400,

            detail=(
                "Unable to extract text from PDF: "
                f"{error}"
            ),

        )


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

async def extract_text_from_docx(
    document: UploadFile
) -> str:

    try:

        file_bytes = await document.read()


        if not file_bytes:

            raise ValueError(
                "DOCX file is empty."
            )


        doc = DocxDocument(
            io.BytesIO(file_bytes)
        )


        paragraphs = []


        # Normal paragraphs

        for paragraph in doc.paragraphs:

            paragraph_text = (
                paragraph.text.strip()
            )


            if paragraph_text:

                paragraphs.append(
                    paragraph_text
                )


        # Tables

        for table in doc.tables:

            for row in table.rows:

                row_text = []


                for cell in row.cells:

                    cell_text = (
                        cell.text.strip()
                    )


                    if cell_text:

                        row_text.append(
                            cell_text
                        )


                if row_text:

                    paragraphs.append(
                        " ".join(row_text)
                    )


        return "\n".join(
            paragraphs
        ).strip()


    except Exception as error:

        logger.exception(
            "DOCX extraction failed: %s",
            error
        )


        raise HTTPException(

            status_code=400,

            detail=(
                "Unable to extract text from DOCX: "
                f"{error}"
            ),

        )


# ============================================================
# TXT EXTRACTION
# ============================================================

async def extract_text_from_txt(
    document: UploadFile
) -> str:

    try:

        file_bytes = await document.read()


        if not file_bytes:

            raise ValueError(
                "TXT file is empty."
            )


        try:

            text = file_bytes.decode(
                "utf-8"
            )

        except UnicodeDecodeError:

            text = file_bytes.decode(
                "latin-1"
            )


        return text.strip()


    except Exception as error:

        logger.exception(
            "TXT extraction failed: %s",
            error
        )


        raise HTTPException(

            status_code=400,

            detail=(
                "Unable to read TXT file: "
                f"{error}"
            ),

        )


# ============================================================
# IMAGE OCR
# ============================================================

async def extract_text_from_image(
    document: UploadFile
) -> str:

    if ocr_reader is None:

        raise RuntimeError(
            "OCR model is not loaded."
        )


    try:

        file_bytes = await document.read()


        if not file_bytes:

            raise ValueError(
                "Image file is empty."
            )


        # Validate image

        image = Image.open(
            io.BytesIO(file_bytes)
        )

        image.load()


        # EasyOCR

        results = ocr_reader.readtext(
            file_bytes,
            detail=0
        )


        return " ".join(
            results
        ).strip()


    except Exception as error:

        logger.exception(
            "Image OCR failed: %s",
            error
        )


        raise HTTPException(

            status_code=400,

            detail=(
                "Unable to process image: "
                f"{error}"
            ),

        )


# ============================================================
# DOCUMENT EXTRACTION ROUTER
# ============================================================

async def extract_document_text(
    document: UploadFile
) -> str:

    filename = (
        document.filename
        or ""
    )


    extension = os.path.splitext(
        filename
    )[1].lower()


    # Validate extension

    if extension not in SUPPORTED_EXTENSIONS:

        raise HTTPException(

            status_code=400,

            detail=(
                "Unsupported document type. "
                "Supported formats: PDF, DOCX, TXT, "
                "JPG, JPEG, PNG and WEBP."
            ),

        )


    logger.info(
        "Processing file: %s",
        filename
    )


    if extension == ".pdf":

        return await extract_text_from_pdf(
            document
        )


    if extension == ".docx":

        return await extract_text_from_docx(
            document
        )


    if extension == ".txt":

        return await extract_text_from_txt(
            document
        )


    if extension in {
        ".jpg",
        ".jpeg",
        ".png",
        ".webp",
    }:

        return await extract_text_from_image(
            document
        )


    return ""


# ============================================================
# CREATE LONG DOCUMENT CHUNKS
# ============================================================

def create_text_chunks(
    text: str,
    max_length: int = MAX_LENGTH,
    stride: int = CHUNK_STRIDE,
) -> list:

    """
    Automatically divides long documents into
    overlapping DistilBERT-compatible chunks.

    Example:

    10,000 tokens

          ↓

    Chunk 1: 0 - 512
    Chunk 2: 448 - 960
    Chunk 3: 896 - 1408
    ...
    """

    if tokenizer is None:

        raise RuntimeError(
            "Tokenizer is not loaded."
        )


    if not text:

        return []


    encoded = tokenizer(

        text,

        add_special_tokens=True,

        truncation=True,

        max_length=max_length,

        stride=stride,

        return_overflowing_tokens=True,

        return_offsets_mapping=True,

        padding=False,

    )


    chunks = []


    input_ids_list = (
        encoded["input_ids"]
    )


    offset_mapping_list = (
        encoded["offset_mapping"]
    )


    for index in range(
        len(input_ids_list)
    ):

        chunks.append({

            "input_ids": (
                input_ids_list[index]
            ),

            "attention_mask": (
                encoded[
                    "attention_mask"
                ][index]
            ),

            "offset_mapping": (
                offset_mapping_list[index]
            ),

        })


    return chunks


# ============================================================
# NER ON ONE CHUNK
# ============================================================

def extract_entities_from_chunk(
    text: str,
    chunk: dict
) -> list:

    if ner_model is None:

        raise RuntimeError(
            "NER model is not loaded."
        )


    # Convert to tensors

    input_ids = torch.tensor(

        [chunk["input_ids"]],

        dtype=torch.long

    )


    attention_mask = torch.tensor(

        [chunk["attention_mask"]],

        dtype=torch.long

    )


    # Run model

    with torch.no_grad():

        outputs = ner_model(

            input_ids=input_ids,

            attention_mask=attention_mask,

        )


    predictions = torch.argmax(

        outputs.logits,

        dim=2

    )[0]


    id2label = (
        ner_model.config.id2label
    )


    offsets = (
        chunk["offset_mapping"]
    )


    entities = []


    # Map predictions to text

    for index, prediction in enumerate(
        predictions
    ):

        label = id2label.get(

            prediction.item(),

            "O"

        )


        if label == "O":

            continue


        if index >= len(offsets):

            continue


        start, end = offsets[index]


        # Ignore special tokens

        if start == end:

            continue


        entity_text = text[
            start:end
        ].strip()


        if not entity_text:

            continue


        entities.append({

            "text": entity_text,

            "label": label,

            "start": start,

            "end": end,

        })


    return entities


# ============================================================
# RUN DISTILBERT NER
# ============================================================

def run_distilbert_ner(
    text: str
) -> list:

    if tokenizer is None:

        raise RuntimeError(
            "DistilBERT tokenizer is not loaded."
        )


    if ner_model is None:

        raise RuntimeError(
            "DistilBERT NER model is not loaded."
        )


    if not text:

        return []


    # Create chunks

    chunks = create_text_chunks(

        text,

        max_length=MAX_LENGTH,

        stride=CHUNK_STRIDE,

    )


    logger.info(

        "Document divided into %d NER chunks",

        len(chunks)

    )


    all_entities = []


    # Process each chunk

    for chunk_index, chunk in enumerate(
        chunks
    ):

        logger.info(

            "Processing NER chunk %d/%d",

            chunk_index + 1,

            len(chunks)

        )


        chunk_entities = (
            extract_entities_from_chunk(

                text,

                chunk

            )
        )


        all_entities.extend(
            chunk_entities
        )


    return all_entities


# ============================================================
# MERGE BIO ENTITIES
# ============================================================

def merge_entities(
    entities: list
) -> list:

    if not entities:

        return []


    # Sort according to position

    entities = sorted(

        entities,

        key=lambda entity: (

            entity["start"],

            entity["end"]

        ),

    )


    merged = []

    current = None


    for entity in entities:

        label = entity[
            "label"
        ]

        text = entity[
            "text"
        ]


        # ----------------------------------------------------
        # B-ENTITY
        # ----------------------------------------------------

        if label.startswith("B-"):

            if current is not None:

                merged.append(
                    current
                )


            current = {

                "text": text,

                "label": label[2:],

                "start": entity["start"],

                "end": entity["end"],

            }


        # ----------------------------------------------------
        # I-ENTITY
        # ----------------------------------------------------

        elif label.startswith("I-"):

            entity_type = label[2:]


            if (

                current is not None

                and current["label"]
                == entity_type

                and entity["start"]
                <= current["end"] + 2

            ):

                current["text"] = (

                    current["text"]

                    + " "

                    + text

                )


                current["end"] = (
                    entity["end"]
                )


            else:

                if current is not None:

                    merged.append(
                        current
                    )


                current = {

                    "text": text,

                    "label": entity_type,

                    "start": entity["start"],

                    "end": entity["end"],

                }


        # ----------------------------------------------------
        # UNKNOWN
        # ----------------------------------------------------

        else:

            if current is not None:

                merged.append(
                    current
                )

                current = None


    if current is not None:

        merged.append(
            current
        )


    return merged


# ============================================================
# REMOVE DUPLICATE ENTITIES
# ============================================================

def deduplicate_entities(
    entities: list
) -> list:

    """
    Removes duplicates created by overlapping chunks.
    """

    unique_entities = []

    seen = set()


    for entity in entities:

        text = (
            entity["text"]
            .strip()
        )


        label = (
            entity["label"]
            .upper()
        )


        key = (

            label,

            text.lower(),

        )


        if key in seen:

            continue


        seen.add(key)


        unique_entities.append(
            entity
        )


    return unique_entities


# ============================================================
# FORMAT SIF ENTITIES
# ============================================================

def format_sif_entities(
    entities: list
) -> dict:

    result = {

        "hazards": [],

        "activities": [],

        "locations": [],

        "barrier_failures": [],

    }


    label_mapping = {

        "HAZARD": "hazards",

        "ACTIVITY": "activities",

        "LOCATION": "locations",

        "BARRIER_FAILURE": (
            "barrier_failures"
        ),

    }


    for entity in entities:

        category = label_mapping.get(

            entity["label"].upper()

        )


        if category is None:

            continue


        value = (
            entity["text"].strip()
        )


        if not value:

            continue


        existing_values = [

            existing.lower()

            for existing
            in result[category]

        ]


        if value.lower() not in existing_values:

            result[category].append(
                value
            )


    return result


# ============================================================
# SIF PRECURSOR SEVERITY CLASSIFICATION
# ============================================================

def classify_sif_precursor_severity(
    entities: dict
) -> dict:

    """
    Rule-based SIF precursor severity classifier.

    Severity scale:

        1 = Very Low
        2 = Low
        3 = Moderate
        4 = High
        5 = Critical

    IMPORTANT:
    This is an initial prototype classifier.

    It should eventually be replaced or calibrated
    using a labelled SIF severity dataset.
    """

    hazards = [
        str(x).lower()
        for x in entities.get(
            "hazards",
            []
        )
    ]


    activities = [
        str(x).lower()
        for x in entities.get(
            "activities",
            []
        )
    ]


    locations = [
        str(x).lower()
        for x in entities.get(
            "locations",
            []
        )
    ]


    barrier_failures = [
        str(x).lower()
        for x in entities.get(
            "barrier_failures",
            []
        )
    ]


    # ========================================================
    # Base severity
    # ========================================================

    score = 1


    # ========================================================
    # HIGH-SEVERITY HAZARDS
    # ========================================================

    critical_hazard_keywords = [

        "fire",

        "explosion",

        "electrocution",

        "electric shock",

        "live wire",

        "high voltage",

        "fall",

        "fall from height",

        "working at height",

        "confined space",

        "toxic gas",

        "poisonous gas",

        "gas leak",

        "chemical exposure",

        "structural collapse",

        "collapse",

        "crane",

        "heavy load",

        "fatal",

    ]


    high_hazard_keywords = [

        "sparks",

        "flame",

        "hot surface",

        "burn",

        "pressure",

        "rotating machinery",

        "moving machinery",

        "machine",

        "cutting",

        "crushing",

        "entanglement",

        "vehicle",

        "forklift",

        "dust",

        "welding",

    ]


    # ========================================================
    # HIGH-SEVERITY ACTIVITIES
    # ========================================================

    critical_activity_keywords = [

        "working at height",

        "confined space entry",

        "electrical maintenance",

        "live electrical work",

        "lifting heavy load",

        "crane operation",

        "demolition",

        "excavation",

    ]


    high_activity_keywords = [

        "welding",

        "cutting",

        "grinding",

        "drilling",

        "maintenance",

        "lifting",

        "construction",

        "scaffolding",

    ]


    # ========================================================
    # CRITICAL BARRIER FAILURES
    # ========================================================

    critical_barrier_keywords = [

        "not wearing ppe",

        "no ppe",

        "without ppe",

        "no fall protection",

        "without fall protection",

        "fall protection missing",

        "guard removed",

        "machine guard removed",

        "safety guard removed",

        "lockout tagout not followed",

        "loto not followed",

        "no isolation",

        "electrical isolation missing",

        "working on live equipment",

        "bypass safety interlock",

        "safety interlock bypassed",

    ]


    high_barrier_keywords = [

        "ppe missing",

        "helmet missing",

        "gloves missing",

        "safety shoes missing",

        "no barricade",

        "barricade missing",

        "warning sign missing",

        "unsafe procedure",

        "procedure not followed",

        "inspection not completed",

        "permit not available",

        "permit missing",

    ]


    # ========================================================
    # Keyword matching
    # ========================================================

    critical_hazard_found = any(

        keyword in hazard

        for hazard in hazards

        for keyword
        in critical_hazard_keywords

    )


    high_hazard_found = any(

        keyword in hazard

        for hazard in hazards

        for keyword
        in high_hazard_keywords

    )


    critical_activity_found = any(

        keyword in activity

        for activity in activities

        for keyword
        in critical_activity_keywords

    )


    high_activity_found = any(

        keyword in activity

        for activity in activities

        for keyword
        in high_activity_keywords

    )


    critical_barrier_found = any(

        keyword in failure

        for failure in barrier_failures

        for keyword
        in critical_barrier_keywords

    )


    high_barrier_found = any(

        keyword in failure

        for failure in barrier_failures

        for keyword
        in high_barrier_keywords

    )


    # ========================================================
    # Calculate severity
    # ========================================================

    # Critical hazard
    if critical_hazard_found:

        score = max(
            score,
            5
        )


    # Critical activity
    if critical_activity_found:

        score = max(
            score,
            5
        )


    # Critical barrier failure
    if critical_barrier_found:

        score = max(
            score,
            5
        )


    # High hazard
    if high_hazard_found:

        score = max(
            score,
            4
        )


    # High-risk activity
    if high_activity_found:

        score = max(
            score,
            4
        )


    # High barrier failure
    if high_barrier_found:

        score = max(
            score,
            4
        )


    # ========================================================
    # Combination rules
    # ========================================================

    # Hazard + barrier failure
    if hazards and barrier_failures:

        score = max(
            score,
            4
        )


    # Multiple hazards
    if len(hazards) >= 3:

        score = max(
            score,
            4
        )


    # Multiple hazards + barrier failure
    if (
        len(hazards) >= 2
        and barrier_failures
    ):

        score = max(
            score,
            5
        )


    # Critical hazard + barrier failure
    if (
        critical_hazard_found
        and barrier_failures
    ):

        score = 5


    # Critical activity + barrier failure
    if (
        critical_activity_found
        and barrier_failures
    ):

        score = 5


    # ========================================================
    # Determine level
    # ========================================================

    severity_levels = {

        1: "Very Low",

        2: "Low",

        3: "Moderate",

        4: "High",

        5: "Critical",

    }


    severity_level = (
        severity_levels[score]
    )


    # ========================================================
    # Explanation
    # ========================================================

    reasons = []


    if critical_hazard_found:

        reasons.append(
            "Critical-severity hazard detected"
        )


    elif high_hazard_found:

        reasons.append(
            "High-severity hazard detected"
        )


    if critical_activity_found:

        reasons.append(
            "High-consequence activity detected"
        )


    elif high_activity_found:

        reasons.append(
            "Higher-risk activity detected"
        )


    if critical_barrier_found:

        reasons.append(
            "Critical barrier failure detected"
        )


    elif high_barrier_found:

        reasons.append(
            "Safety barrier failure detected"
        )


    if len(hazards) >= 2:

        reasons.append(
            "Multiple hazards detected"
        )


    if not reasons:

        reasons.append(
            "No major high-consequence precursor "
            "indicators detected"
        )


    return {

        "score": score,

        "level": severity_level,

        "scale": "1-5",

        "reasons": reasons,

    }


# ============================================================
# COMPLETE REPORT PROCESSING
# ============================================================

def process_report_text(
    raw_text: str
) -> dict:

    # --------------------------------------------------------
    # Clean
    # --------------------------------------------------------

    cleaned_text = clean_report_text(
        raw_text
    )


    if not cleaned_text:

        raise ValueError(
            "No readable text was found."
        )


    # --------------------------------------------------------
    # DistilBERT NER
    # --------------------------------------------------------

    token_entities = run_distilbert_ner(
        cleaned_text
    )


    # --------------------------------------------------------
    # Merge BIO
    # --------------------------------------------------------

    merged_entities = merge_entities(
        token_entities
    )


    # --------------------------------------------------------
    # Remove duplicates
    # --------------------------------------------------------

    unique_entities = (
        deduplicate_entities(
            merged_entities
        )
    )


    # --------------------------------------------------------
    # Format SIF entities
    # --------------------------------------------------------

    sif_entities = format_sif_entities(
        unique_entities
    )


    # --------------------------------------------------------
    # SIF precursor severity
    # --------------------------------------------------------

    severity = (
        classify_sif_precursor_severity(
            sif_entities
        )
    )


    return {

        "cleaned_text": cleaned_text,

        "entities": sif_entities,

        "raw_ner_entities": (
            unique_entities
        ),

        "sif_precursor_severity": severity,

    }


# ============================================================
# PROCESS DOCUMENT ENDPOINT
# ============================================================

@app.post("/process-document")
async def process_document(

    document: Optional[
        UploadFile
    ] = File(None),

    report_text: Optional[
        str
    ] = Form(None),

):

    try:

        # ----------------------------------------------------
        # Validate input
        # ----------------------------------------------------

        if (
            document is None
            and not report_text
        ):

            raise HTTPException(

                status_code=400,

                detail=(
                    "Please provide either "
                    "a document or report_text."
                ),

            )


        raw_text = ""


        # ====================================================
        # FILE PROCESSING
        # ====================================================

        if document is not None:

            file_bytes = await document.read()


            if not file_bytes:

                raise HTTPException(

                    status_code=400,

                    detail=(
                        "Uploaded file is empty."
                    ),

                )


            # File size

            if len(file_bytes) > MAX_FILE_SIZE:

                raise HTTPException(

                    status_code=413,

                    detail=(
                        "File size exceeds "
                        "the 10 MB limit."
                    ),

                )


            # Reset pointer

            await document.seek(0)


            # Extract text

            raw_text = (
                await extract_document_text(
                    document
                )
            )


        # ====================================================
        # DIRECT REPORT TEXT
        # ====================================================

        if report_text:

            report_text = (
                report_text.strip()
            )


            if report_text:

                if raw_text:

                    raw_text = (

                        raw_text

                        + "\n"

                        + report_text

                    )

                else:

                    raw_text = report_text


        # ====================================================
        # VALIDATE EXTRACTED TEXT
        # ====================================================

        if not raw_text.strip():

            raise HTTPException(

                status_code=400,

                detail=(
                    "No readable text was found "
                    "in the submitted report."
                ),

            )


        # ====================================================
        # AI PROCESSING
        # ====================================================

        logger.info(
            "Starting SIF text processing..."
        )


        result = process_report_text(
            raw_text
        )


        # ====================================================
        # RISK ENGINE
        # ====================================================

        logger.info(
            "Running SIF risk engine..."
        )


        risk_assessment = (
            risk_engine.assess_risk(
                result["entities"]
            )
        )


        logger.info(
            "SIF processing completed successfully."
        )


        # ====================================================
        # RESPONSE
        # ====================================================

        return {

            "status": "success",


            "processing": {

                "document_received": (
                    document is not None
                ),

                "text_cleaning": True,

                "long_document_chunking": True,

                "chunk_size": MAX_LENGTH,

                "chunk_overlap": CHUNK_STRIDE,

                "distilbert_ner": True,

                "sif_precursor_severity": True,

                "risk_engine": True,

            },


            # ------------------------------------------------
            # Raw extracted text
            # ------------------------------------------------

            "raw_text": raw_text,


            # ------------------------------------------------
            # Cleaned text
            # ------------------------------------------------

            "cleaned_text": (
                result["cleaned_text"]
            ),


            # ------------------------------------------------
            # SIF entities
            # ------------------------------------------------

            "entities": result[
                "entities"
            ],


            # ------------------------------------------------
            # Raw NER output
            # ------------------------------------------------

            "raw_ner_entities": result[
                "raw_ner_entities"
            ],


            # ------------------------------------------------
            # SIF PRECURSOR SEVERITY
            # ------------------------------------------------

            "sif_precursor_severity": (
                result[
                    "sif_precursor_severity"
                ]
            ),


            # ------------------------------------------------
            # Existing risk engine
            # ------------------------------------------------

            "risk_assessment": (
                risk_assessment
            ),

        }


    except HTTPException:

        raise


    except Exception as error:

        logger.exception(
            "SIF processing failed: %s",
            error
        )


        raise HTTPException(

            status_code=500,

            detail=(
                "SIF document processing failed."
            ),

        )


# ============================================================
# MANUAL RISK ASSESSMENT ENDPOINT
# ============================================================

@app.post("/assess-risk")
async def assess_risk_endpoint(
    request: RiskAssessmentRequest
):

    entities = {

        "hazards": request.hazards,

        "activities": request.activities,

        "locations": request.locations,

        "barrier_failures": (
            request.barrier_failures
        ),

    }


    if not any(
        entities.values()
    ):

        raise HTTPException(

            status_code=400,

            detail=(
                "At least one entity list "
                "must be non-empty."
            ),

        )


    # --------------------------------------------------------
    # SIF precursor severity
    # --------------------------------------------------------

    severity = (
        classify_sif_precursor_severity(
            entities
        )
    )


    # --------------------------------------------------------
    # Existing risk engine
    # --------------------------------------------------------

    risk_assessment = (
        risk_engine.assess_risk(
            entities
        )
    )


    return {

        "status": "success",

        "sif_precursor_severity": severity,

        "risk_assessment": (
            risk_assessment
        ),

    }


# ============================================================
# SERVER START
# ============================================================

if __name__ == "__main__":

    import uvicorn


    uvicorn.run(

        "main:app",

        host="127.0.0.1",

        port=8000,

        reload=True,

    )
